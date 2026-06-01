import asyncio
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path

import requests
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from copilot import CopilotClient, SubprocessConfig
from copilot.session import PermissionHandler


# -------------------------
# ENV
# -------------------------
GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]
COPILOT_GITHUB_TOKEN = os.environ["COPILOT_GITHUB_TOKEN"]
REPO = os.environ["GITHUB_REPOSITORY"]
BRANCH = os.environ["TARGET_BRANCH"]
START_DATE = os.environ["START_DATE"]
END_DATE = os.environ["END_DATE"]

OUTPUT_DIR = Path("output")
FONT_NAME = "Times New Roman"
HIGHLIGHT_YELLOW = "FFF200"


# -------------------------
# STYLE HELPERS
# -------------------------
def apply_font(run, size=None, bold=None, italic=None, color=None, underline=None):
    run.font.name = FONT_NAME
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if underline is not None:
        run.underline = underline


def set_document_font(doc):
    for style_name, size in [
        ("Normal", 12),
        ("Title", 24),
        ("Heading 1", 16),
        ("Heading 2", 14),
        ("Heading 3", 12),
    ]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_NAME
            style.font.size = Pt(size)


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_paragraph_with_font(doc, text, size=12, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    apply_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_with_font(doc, text, level=1, highlight=False):
    p = doc.add_paragraph()
    if level == 1:
        p.style = doc.styles["Heading 1"]
        size = 16
    elif level == 2:
        p.style = doc.styles["Heading 2"]
        size = 14
    else:
        p.style = doc.styles["Heading 3"]
        size = 12

    run = p.add_run(text)
    apply_font(run, size=size, bold=True)

    if highlight:
        for r in p.runs:
            r.font.highlight_color = 7  # yellow highlight in Word
    return p


def format_cell(cell, text="", size=10, bold=False, align=WD_PARAGRAPH_ALIGNMENT.LEFT, default="N/A"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    value = text if text not in [None, ""] else default
    if value != "":
        run = p.add_run(value)
        apply_font(run, size=size, bold=bold)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_header_cell(cell, text, size=10):
    format_cell(cell, text, size=size, bold=True, align=WD_PARAGRAPH_ALIGNMENT.CENTER, default="")
    shade_cell(cell, HIGHLIGHT_YELLOW)


def add_email_line(doc, name, email, size=12):
    p = doc.add_paragraph()
    run1 = p.add_run(name + " ")
    apply_font(run1, size=size, bold=True)

    run2 = p.add_run(email)
    apply_font(run2, size=size, bold=True, color=RGBColor(0, 0, 255), underline=True)


def add_table(doc, headers, rows, col_widths, header_size=10, body_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for i, width in enumerate(col_widths):
        table.columns[i].width = Inches(width)
        for cell in table.columns[i].cells:
            cell.width = Inches(width)

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_header_cell(hdr_cells[i], header, size=header_size)

    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            format_cell(row[i], value, size=body_size, default="")

    return table


# -------------------------
# DATE FILTER
# -------------------------
def is_within_range(merged_at):
    if not merged_at:
        return False

    merged_dt = datetime.fromisoformat(merged_at.replace("Z", "+00:00"))
    start_dt = datetime.fromisoformat(START_DATE).replace(tzinfo=timezone.utc)
    end_dt = datetime.fromisoformat(END_DATE + "T23:59:59").replace(tzinfo=timezone.utc)
    return start_dt <= merged_dt <= end_dt


# -------------------------
# FETCH PRs
# -------------------------
def fetch_prs():
    url = f"https://api.github.com/repos/{REPO}/pulls"
    headers = {
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }

    prs = []
    page = 1

    while True:
        params = {
            "state": "closed",
            "base": BRANCH,
            "page": page,
            "per_page": 100,
            "sort": "updated",
            "direction": "desc",
        }

        r = requests.get(url, headers=headers, params=params, timeout=60)
        if r.status_code != 200:
            raise RuntimeError(f"GitHub API failed: {r.status_code} {r.text}")

        batch = r.json()
        if not batch:
            break

        for pr in batch:
            if pr.get("merged_at") and is_within_range(pr["merged_at"]):
                prs.append(pr)

        if len(batch) < 100:
            break
        page += 1

    return prs


def split_us_defect_and_description(text):
    """
    Examples:
      SPM/..._Feature abc  -> US/Defect# = SPM/..., Description = Feature abc
      SPER/..._Update xyz   -> US/Defect# = SPER/..., Description = Update xyz

    Rule:
      everything before first underscore = US/Defect#
      everything after first underscore  = Description
    """
    value = " ".join((text or "").split()).strip()
    if "_" in value:
        left, right = value.split("_", 1)
        return left.strip(), right.strip()
    return "", value


def prepare_data(prs):
    result = []
    for pr in prs:
        title = pr.get("title", "") or ""
        us_defect_no, description = split_us_defect_and_description(title)

        jira = re.findall(r"[A-Z]+-\d+", title)
        merged_at = pr.get("merged_at", "")

        result.append(
            {
                "track_id": ", ".join(jira) if jira else f"PR-{pr['number']}",
                "us_defect_no": us_defect_no,
                "description": description or title,
                "author": (pr.get("user") or {}).get("login", ""),
                "merged_at": format_merge_date(merged_at),
                "url": pr.get("html_url", ""),
            }
        )
    return result


def format_merge_date(merged_at):
    try:
        dt = datetime.fromisoformat(merged_at.replace("Z", "+00:00"))
        return dt.strftime("%Y-%m-%d %H:%M UTC")
    except Exception:
        return merged_at or "N/A"


# -------------------------
# COPILOT
# -------------------------
def extract_json(text):
    text = (text or "").strip()
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except json.JSONDecodeError:
            pass

    return {
        "introduction": "Release changes included.",
        "dependencies": "No major dependencies identified.",
    }


async def generate_ai(data):
    client = CopilotClient(
        SubprocessConfig(
            github_token=COPILOT_GITHUB_TOKEN,
            use_logged_in_user=False,
        )
    )

    async with client:
        session = await client.create_session(
            model="gpt-4.1",
            session_id="release-summary",
            on_permission_request=PermissionHandler.approve_all,
        )

        prompt = f"""
Generate concise enterprise release-note JSON only.

Return exactly:
{{
  "introduction": "...",
  "dependencies": "..."
}}

Use professional language for release notes.
Do not add markdown.

Important parsing rule for the PR data:
- If a title/description starts with SPM, SPER, or SPNR and contains an underscore `_`,
  treat everything before the first `_` as the US/Defect#.
- Treat everything after the first `_` as the Description.

PR Data:
{json.dumps(data, indent=2)}
"""

        resp = await session.send_and_wait(prompt)
        return extract_json(resp.data.content)


# -------------------------
# WORD DOCUMENT
# -------------------------
def create_doc(ai, data):
    OUTPUT_DIR.mkdir(exist_ok=True)

    doc = Document()
    set_document_font(doc)

    section = doc.sections[0]
    set_landscape(section)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    branch_display = BRANCH.replace("-", " ").replace("_", " ").title()

    run = title.add_run(f"Release Management\nDocument - Release Notes\nfor {branch_display} Branch")
    apply_font(run, size=24, bold=False)

    doc.add_paragraph("")

    # 1. Introduction
    add_heading_with_font(doc, "1. Introduction", level=1)
    intro_text = ai.get(
        "introduction",
        "Will be created for a Sprint and will include all user stories and corresponding packages that will be merged to Dev Master."
    )
    add_paragraph_with_font(doc, intro_text, size=12)

    # 2. Dependencies
    add_heading_with_font(doc, "2. Dependencies", level=1)
    dep_text = ai.get(
        "dependencies",
        "This release depends on successful integration of updated SQL views and downstream validation by the data engineering and QA teams."
    )
    add_paragraph_with_font(doc, dep_text, size=12)

    # 3. Implementation Summary
    add_heading_with_font(doc, "3. Implementation Summary", level=1)

    headers = ["Track ID", "US/Defect#", "Description", "Author", "Merge Date", "Reference"]
    rows = [
        [pr["track_id"], pr["us_defect_no"], pr["description"], pr["author"], pr["merged_at"], pr["url"]]
        for pr in data
    ]
    col_widths = [1.05, 1.75, 3.55, 1.2, 1.45, 2.05]
    add_table(doc, headers, rows, col_widths, header_size=10, body_size=9)

    # 4. Danone PO
    add_heading_with_font(doc, "4. Danone PO", level=1, highlight=True)
    danone_po = [
        ("CHICOULAA Grégoire", "Gregoire.CHICOULAA@danone.com"),
        ("DE MOEGEN Cyril", "Cyril.DE-MOEGEN@danone.com"),
        ("PULIDO Catalina", "Catalina.PULIDO@danone.com"),
        ("HAUPTMANN Clemens", "Clemens.HAUPTMANN@danone.com"),
        ("ZELAZEK Lukasz", "Lukasz.ZELAZEK1@danone.com"),
        ("BOOTH Christopher", "Christopher.BOOTH@danone.com"),
        ("ZHURAVLEVA Antonina (EXT)", "Antonina.ZHURAVLEVA@external.danone.com"),
        ("GAILLY Eric (EXT)", "Eric.GAILLY@external.danone.com"),
        ("PECLAK Justyna", "Justyna.PECLAK@danone.com"),
        ("DE KRUIJK Harry", "Harry.DEKRUIJK@danone.com"),
        ("SZYMCZYK Radoslaw", "Radoslaw.SZYMCZYK@danone.com"),
    ]
    for name, email in danone_po:
        add_email_line(doc, name, email, size=12)

    # 5. 09 Team
    add_heading_with_font(doc, "5. 09 Team", level=1, highlight=True)
    team_09 = [
        ("Rizvana Shaik", "rizvana.shaik@09solutions.net"),
        ("Rashid Patel", "rashid.patel@09solutions.com"),
        ("Sunny Singh", "sunny.singh@09solutions.com"),
        ("Sumit Kumar", "sumit.kumar@09solutions.com"),
    ]
    for name, email in team_09:
        add_email_line(doc, name, email, size=12)

    # 7. Approvals
    add_heading_with_font(doc, "7. Approvals", level=1, highlight=True)
    approvals_headers = ["Role", "Name", "Date", "User Story", "Approved"]
    approvals_rows = [
        ["Release Manager", "", "", "", "Yes"],
        ["Stakeholder", "", "", "", "Yes"],
    ]
    approvals_widths = [2.6, 1.9, 1.3, 2.0, 1.3]
    add_table(doc, approvals_headers, approvals_rows, approvals_widths, header_size=10, body_size=10)

    # 8. Internal Team Approvals
    add_heading_with_font(doc, "8. Internal Team Approvals", level=1, highlight=True)
    internal_headers = ["US/Defect#", "Role", "Name", "Date", "Approved"]
    internal_rows = [
        ["", "Test View Sign Off by PO", "", "", "Yes"],
        ["", "Infosys Lead Sign Off", "", "", "Yes"],
        ["", "D&A stakeholder Sign off", "", "", "Yes"],
    ]
    internal_widths = [1.4, 4.0, 1.8, 1.4, 1.2]
    add_table(doc, internal_headers, internal_rows, internal_widths, header_size=10, body_size=10)

    # 9. Testing Results
    add_heading_with_font(doc, "9. Testing Results", level=1, highlight=True)
    testing_headers = ["US/Defect#", "Sprint #", "BGL", "Owner", "Test Results Link"]
    testing_rows = [
        ["", "", "", "", ""],
        ["", "", "", "", ""],
    ]
    testing_widths = [1.5, 1.2, 1.2, 1.8, 3.5]
    add_table(doc, testing_headers, testing_rows, testing_widths, header_size=10, body_size=10)

    file_path = OUTPUT_DIR / f"pr-summary-{BRANCH}.docx"
    doc.save(file_path)
    print(f"Saved: {file_path}")


# -------------------------
# MAIN
# -------------------------
async def main():
    print("Fetching PRs...")
    prs = fetch_prs()
    print(f"Filtered PRs: {len(prs)}")

    data = prepare_data(prs)

    if not data:
        create_doc(
            {
                "introduction": "No PRs found in the selected date range.",
                "dependencies": "N/A",
            },
            data,
        )
        return

    print("Generating AI content...")
    ai = await generate_ai(data)

    print("Creating document...")
    create_doc(ai, data)


if __name__ == "__main__":
    asyncio.run(main())
