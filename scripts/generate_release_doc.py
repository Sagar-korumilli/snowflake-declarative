#!/usr/bin/env python3
import argparse
import datetime as dt
import difflib
import os
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


@dataclass
class FileChange:
    status: str
    old_path: Optional[str]
    new_path: Optional[str]
    insertions: Optional[int] = None
    deletions: Optional[int] = None


def run(cmd: List[str]) -> str:
    result = subprocess.run(cmd, capture_output=True, text=True, check=True)
    return result.stdout.strip()


def git_show(rev: str, path: str) -> Optional[str]:
    try:
        out = subprocess.run(
            ["git", "show", f"{rev}:{path}"],
            capture_output=True,
            check=True,
        )
        return out.stdout.decode("utf-8", errors="replace")
    except subprocess.CalledProcessError:
        return None


def git_numstat(base_sha: str, head_sha: str, path: str) -> Tuple[Optional[int], Optional[int]]:
    try:
        out = run(["git", "diff", "--numstat", base_sha, head_sha, "--", path])
        if not out:
            return None, None
        parts = out.split("\t")
        if len(parts) >= 3:
            ins = None if parts[0] == "-" else int(parts[0])
            dels = None if parts[1] == "-" else int(parts[1])
            return ins, dels
    except Exception:
        pass
    return None, None


def parse_changed_files(base_sha: str, head_sha: str) -> List[FileChange]:
    out = run(["git", "diff", "--name-status", "-M", base_sha, head_sha, "--"])
    changes: List[FileChange] = []

    if not out:
        return changes

    for line in out.splitlines():
        parts = line.split("\t")
        status = parts[0]

        if status.startswith("R") or status.startswith("C"):
            # Rename/copy: R100 old new
            old_path = parts[1]
            new_path = parts[2]
        elif status == "A":
            old_path = None
            new_path = parts[1]
        elif status == "D":
            old_path = parts[1]
            new_path = None
        else:
            old_path = parts[1]
            new_path = parts[1]

        ins, dels = git_numstat(base_sha, head_sha, new_path or old_path or "")
        changes.append(
            FileChange(
                status=status,
                old_path=old_path,
                new_path=new_path,
                insertions=ins,
                deletions=dels,
            )
        )
    return changes


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9, mono: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"
        r = run._element.rPr
        r.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_table_row(table, values, bold=False, mono=False):
    row = table.add_row().cells
    for i, v in enumerate(values):
        set_cell_text(row[i], str(v), bold=bold, mono=mono)


def add_wrapped_text_paragraph(doc: Document, title: str, text: str, max_chars: int = 6000):
    doc.add_paragraph(title, style="Heading 3")
    if text is None:
        text = "(not available)"
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n...[truncated]..."
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_section_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.add_run(text)


def build_diff(old_text: Optional[str], new_text: Optional[str], old_label: str, new_label: str) -> str:
    old_lines = old_text.splitlines() if old_text is not None else []
    new_lines = new_text.splitlines() if new_text is not None else []
    diff = difflib.unified_diff(
        old_lines,
        new_lines,
        fromfile=old_label,
        tofile=new_label,
        lineterm="",
        n=3,
    )
    return "\n".join(diff)


def create_doc(
    output: Path,
    pr_number: str,
    pr_title: str,
    author: str,
    base_branch: str,
    head_branch: str,
    base_sha: str,
    head_sha: str,
    repo: str,
    changes: List[FileChange],
):
    doc = Document()

    # Margins
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # Default font
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Release Management Document")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run(f"PR #{pr_number} | {repo}")
    r.italic = True
    r.font.size = Pt(10)

    doc.add_paragraph("")

    # Overview table
    doc.add_paragraph("Pull Request Overview", style="Heading 1")
    overview = doc.add_table(rows=0, cols=2)
    overview.style = "Table Grid"
    overview.autofit = False
    overview.columns[0].width = Inches(2.0)
    overview.columns[1].width = Inches(4.9)

    overview_rows = [
        ("PR Number", pr_number),
        ("Title", pr_title),
        ("Author", author),
        ("Source Branch", head_branch),
        ("Target Branch", base_branch),
        ("Base Commit", base_sha),
        ("Head Commit", head_sha),
        ("Generated At", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for k, v in overview_rows:
        row = overview.add_row().cells
        set_cell_text(row[0], k, bold=True, size=9)
        set_cell_text(row[1], v, size=9)

    doc.add_paragraph("")

    # Changed files summary
    add_section_heading(doc, "Changed Files Summary")
    summary = doc.add_table(rows=1, cols=5)
    summary.style = "Table Grid"
    summary.autofit = False
    widths = [0.6, 0.9, 3.6, 1.0, 1.0]
    for i, w in enumerate(widths):
        summary.columns[i].width = Inches(w)

    hdr = summary.rows[0].cells
    headers = ["#", "Status", "File", "+", "-"]
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, size=9)

    for idx, ch in enumerate(changes, start=1):
        path_display = ch.new_path or ch.old_path or ""
        add_table_row(
            summary,
            [
                idx,
                ch.status,
                path_display,
                ch.insertions if ch.insertions is not None else "",
                ch.deletions if ch.deletions is not None else "",
            ],
            bold=False,
        )

    doc.add_paragraph("")

    # Per-file detail sections
    for idx, ch in enumerate(changes, start=1):
        old_path = ch.old_path
        new_path = ch.new_path
        display_path = new_path or old_path or "Unknown file"

        add_section_heading(doc, f"{idx}. {display_path}")

        meta = doc.add_table(rows=0, cols=2)
        meta.style = "Table Grid"
        meta.autofit = False
        meta.columns[0].width = Inches(1.7)
        meta.columns[1].width = Inches(5.2)

        meta_rows = [
            ("Status", ch.status),
            ("Old Path", old_path or "(new file)"),
            ("New Path", new_path or "(deleted file)"),
            ("Insertions", ch.insertions if ch.insertions is not None else ""),
            ("Deletions", ch.deletions if ch.deletions is not None else ""),
        ]
        for k, v in meta_rows:
            row = meta.add_row().cells
            set_cell_text(row[0], k, bold=True, size=9)
            set_cell_text(row[1], str(v), size=9)

        old_text = git_show(base_sha, old_path) if old_path else None
        new_text = git_show(head_sha, new_path) if new_path else None

        doc.add_paragraph("")
        add_wrapped_text_paragraph(doc, "Old File Content", old_text)
        add_wrapped_text_paragraph(doc, "New File Content", new_text)

        diff_text = build_diff(
            old_text,
            new_text,
            old_label=f"{base_branch}/{old_path or display_path}",
            new_label=f"{head_branch}/{new_path or display_path}",
        )
        add_wrapped_text_paragraph(doc, "Unified Diff", diff_text, max_chars=8000)

        doc.add_paragraph("")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="Generate a release management Word document from Git diff.")
    parser.add_argument("--base-sha", required=True)
    parser.add_argument("--head-sha", required=True)
    parser.add_argument("--base-branch", required=True)
    parser.add_argument("--head-branch", required=True)
    parser.add_argument("--pr-number", required=True)
    parser.add_argument("--pr-title", required=True)
    parser.add_argument("--author", required=True)
    parser.add_argument("--repo", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    changes = parse_changed_files(args.base_sha, args.head_sha)
    create_doc(
        output=Path(args.output),
        pr_number=args.pr_number,
        pr_title=args.pr_title,
        author=args.author,
        base_branch=args.base_branch,
        head_branch=args.head_branch,
        base_sha=args.base_sha,
        head_sha=args.head_sha,
        repo=args.repo,
        changes=changes,
    )

    print(f"Created: {args.output}")


if __name__ == "__main__":
    main()
